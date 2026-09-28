"""
Excel Export — چندین Sheet برای بررسی دقیق توسط هیئت‌مدیره/تحلیل‌گر.
"""
from __future__ import annotations
import sqlite3
import sys
import pathlib
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[2]))

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage

from app.analytics import metrics as m
from app.detection import anomaly as an
from app.detection import behavior as bh
from app.analytics import market_maker as mm_mod
from app.analytics import comparison as comp
from app.analytics import insights as ins
from app.analytics import extended_analysis as ext
from app.analytics import portfolio as pf
from app.core.labels import alert_type_fa, severity_fa
from app.core.jalali import jalali_date_display, jalali_datetime_display
from app.config.settings import (get_company_name, get_report_prepared_by, get_report_audience,
                                  get_report_title, get_logo_path, DEVELOPER_CREDIT_LINE)

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(name="Tahoma", color="FFFFFF", bold=True)


def _prepare_sheet(ws):
    """تنظیمات حرفه‌ای و پایدار Excel برای فارسی/RTL و چاپ."""
    ws.sheet_view.rightToLeft = True
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5
    ws.print_options.horizontalCentered = True
    ws.oddFooter.center.text = "صفحه &P از &N"
    ws.oddFooter.center.size = 9
    ws.oddFooter.center.font = "Tahoma"
    ws.oddFooter.right.text = "Smart Equity Transaction Intelligence"
    ws.oddFooter.right.size = 8
    ws.oddFooter.right.font = "Tahoma"
    if ws.max_row and ws.max_column:
        from openpyxl.utils import get_column_letter
        # حداقل عرض عددی کافی برای ارقام ریالی بزرگ؛ از ### در Excel/چاپ جلوگیری می‌کند.
        for c in range(1, ws.max_column + 1):
            letter = get_column_letter(c)
            if ws.column_dimensions[letter].width is None or ws.column_dimensions[letter].width < 18:
                ws.column_dimensions[letter].width = 18
        if ws.max_column >= 2 and ws.title == "خلاصه":
            ws.column_dimensions["A"].width = 30
            ws.column_dimensions["B"].width = 26
        ws.print_area = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"



def _level_fa(level: str) -> str:
    return {"Low concentration": "تمرکز پایین", "Medium concentration": "تمرکز متوسط",
            "High concentration": "تمرکز بالا"}.get(level, level)


def _write_table(ws, headers: list[str], rows: list[list], start_row=1):
    _prepare_sheet(ws)
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col_idx, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r_idx, row in enumerate(rows, start=start_row + 1):
        for c_idx, val in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font = Font(name="Tahoma", size=10)
            cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
            if isinstance(val, float):
                cell.number_format = "#,##0.00"
            elif isinstance(val, int) and not isinstance(val, bool):
                cell.number_format = "#,##0"
    for col_idx in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 22
    ws.sheet_view.rightToLeft = True
    end_row = start_row + len(rows)
    ws.freeze_panes = f"A{start_row + 1}"
    if end_row >= start_row:
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{end_row}"
    return end_row + 2


def _write_narrative(ws, lines: list[list], start_row=1, title=None):
    """می‌نویسد چند خط روایی (هر خط یک ردیف تمام‌عرض با wrap_text) - برای بخش‌های تفسیری."""
    _prepare_sheet(ws)
    row = start_row
    if title:
        cell = ws.cell(row=row, column=1, value=title)
        cell.font = Font(name="Tahoma", bold=True, size=12, color="1F4E78")
        cell.alignment = Alignment(horizontal="right", vertical="center")
        row += 1
    for line in lines:
        cell = ws.cell(row=row, column=1, value=("• " + line) if not line.startswith("•") else line)
        cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
        cell.font = Font(name="Tahoma", size=10)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        row += 1
    return row + 1


def generate_excel_report(conn: sqlite3.Connection, out_path: str, start=None, end=None, top_n=15):
    wb = openpyxl.Workbook()

    # --- Sheet 1: Overview + narrative executive summary ---
    ws = wb.active
    ws.title = "خلاصه"
    ov = m.overview(conn, start, end)
    logo_path = get_logo_path()
    first_content_row = 1
    if logo_path:
        try:
            from app.core.image_utils import fit_dimensions
            w_px, h_px = fit_dimensions(str(logo_path), 160, 70)
            xl_img = XLImage(str(logo_path))
            xl_img.width, xl_img.height = w_px, h_px
            ws.add_image(xl_img, "A1")
            # لوگو به‌صورت شناور از سلول A1 رسم می‌شود و صرفاً روی محدوده استفاده‌شده (ستون‌های ۱-۶)
            # قابل اعتماد است؛ اگر محتوای متنی همان ردیف‌ها را اشغال کند، روی هم می‌افتند و در
            # PDF/چاپ ممکن است لوگو دیده نشود چون خارج از محدوده Print Area تشخیص داده می‌شود.
            # پس چند ردیف بالای صفحه را فقط برای لوگو خالی نگه می‌داریم.
            first_content_row = 5
        except Exception:
            pass  # نبود/خرابی لوگو نباید کل تولید گزارش را متوقف کند
    company_name = get_company_name()
    header_lines = ([f"شرکت: {company_name}"] if company_name else []) + [
        f"تهیه‌کننده: {get_report_prepared_by()}   —   ارائه به: {get_report_audience()}",
    ]
    next_row = _write_narrative(ws, header_lines, start_row=first_content_row, title=get_report_title())
    narrative_lines = ins.executive_summary_narrative(conn, start, end)
    next_row = _write_narrative(ws, narrative_lines, start_row=next_row, title="خلاصه مدیریتی")
    _write_table(ws, ["شاخص", "مقدار"], [
        ["تعداد معاملات", ov["n_transactions"]],
        ["حجم کل", ov["total_quantity"]],
        ["ارزش کل (ریال)", ov["total_value"]],
        ["میانگین وزنی قیمت", round(ov["weighted_avg_price"], 2)],
        ["تعداد خریداران منحصربه‌فرد", ov["n_unique_buyers"]],
        ["تعداد فروشندگان منحصربه‌فرد", ov["n_unique_sellers"]],
        ["بزرگترین معامله (ریال)", ov["largest_transaction_value"]],
        ["از تاریخ", jalali_date_display(ov["first_date"])],
        ["تا تاریخ", jalali_date_display(ov["last_date"])],
    ], start_row=next_row)
    footer_cell = ws.cell(row=next_row + 11, column=1, value=DEVELOPER_CREDIT_LINE)
    footer_cell.font = Font(italic=True, size=9, color="888888")

    # --- Sheet 2: Top buyer brokers ---
    ws2 = wb.create_sheet("کارگزاران خریدار برتر")
    brokers = m.top_brokers(conn, "buyer", start, end, top_n=top_n)
    _write_table(ws2, ["نام کارگزار", "بازارگردان اصلی", "حجم", "ارزش (ریال)", "تعداد معاملات", "تعداد طرف مقابل"],
                 [[b["broker_name"], "بله" if b["is_primary_market_maker"] else "خیر", b["quantity"], b["value"],
                   b["n_transactions"], b["n_counterparties"]] for b in brokers])

    # --- Sheet 3: Top seller brokers ---
    ws3 = wb.create_sheet("کارگزاران فروشنده برتر")
    brokers_s = m.top_brokers(conn, "seller", start, end, top_n=top_n)
    _write_table(ws3, ["نام کارگزار", "بازارگردان اصلی", "حجم", "ارزش (ریال)", "تعداد معاملات", "تعداد طرف مقابل"],
                 [[b["broker_name"], "بله" if b["is_primary_market_maker"] else "خیر", b["quantity"], b["value"],
                   b["n_transactions"], b["n_counterparties"]] for b in brokers_s])

    # --- Sheet 4: Top buyers (persons) ---
    ws4 = wb.create_sheet("خریداران برتر")
    buyers = m.top_persons(conn, "buyer", start, end, top_n=top_n)
    _write_table(ws4, ["نام", "نوع", "اطمینان شناسه", "حجم", "ارزش (ریال)", "تعداد معاملات", "روزهای فعال", "اولین فعالیت", "آخرین فعالیت"],
                 [[p["name"], p["type"], p["id_confidence"], p["quantity"], p["value"], p["n_transactions"],
                   p["active_days"], jalali_date_display(p["first_seen"]), jalali_date_display(p["last_seen"])] for p in buyers])

    # --- Sheet 5: Top sellers (persons) ---
    ws5 = wb.create_sheet("فروشندگان برتر")
    sellers = m.top_persons(conn, "seller", start, end, top_n=top_n)
    _write_table(ws5, ["نام", "نوع", "اطمینان شناسه", "حجم", "ارزش (ریال)", "تعداد معاملات", "روزهای فعال", "اولین فعالیت", "آخرین فعالیت"],
                 [[p["name"], p["type"], p["id_confidence"], p["quantity"], p["value"], p["n_transactions"],
                   p["active_days"], jalali_date_display(p["first_seen"]), jalali_date_display(p["last_seen"])] for p in sellers])

    # --- Sheet 6: Net position (persons) ---
    ws6 = wb.create_sheet("خالص خرید و فروش")
    net = m.net_position_persons(conn, start, end, top_n=top_n)
    rows = [["--- برترین خریداران خالص ---", "", ""]]
    rows += [[p["name"], p["net_value"], p["net_quantity"]] for p in net["top_net_buyers"]]
    rows += [["--- برترین فروشندگان خالص ---", "", ""]]
    rows += [[p["name"], p["net_value"], p["net_quantity"]] for p in net["top_net_sellers"]]
    _write_table(ws6, ["نام", "ارزش خالص (ریال)", "حجم خالص"], rows)

    # --- Sheet 7: Short-term / round-trip traders ---
    ws7 = wb.create_sheet("رفتار کوتاه‌مدت")
    st = bh.top_short_term_traders(conn, start, end, top_n=top_n)
    _write_table(ws7, ["نام", "امتیاز (۰-۱۰۰)", "سطح", "تعداد رفت‌وبرگشت", "روزهای فعال"],
                 [[t["name"], t["score"], t["level"], t["n_round_trips"], t["active_days"]] for t in st])
    ws7.cell(row=len(st) + 3, column=1,
             value="توضیح: این امتیاز صرفاً یک شاخص آماری از الگوی مشاهده‌شده است و اثبات‌کننده قصد یا نوسان‌گیری واقعی نیست.")

    # --- Sheet 8: Anomalies ---
    ws8 = wb.create_sheet("هشدارها")
    alerts = an.run_all_anomaly_detectors(conn, start, end)[:100]
    _write_table(ws8, ["نوع", "شدت", "امتیاز ریسک", "تاریخ", "ارزش (ریال)", "خریدار", "فروشنده", "توضیح"],
                 [[alert_type_fa(a["alert_type"]), severity_fa(a["severity"]), a["risk_score"], jalali_date_display(a.get("trade_date")),
                   a.get("value"), a.get("buyer_name") or "—", a.get("seller_name") or "—", a["reason"]]
                  for a in alerts])

    # --- Sheet 9: Concentration ---
    ws9 = wb.create_sheet("تمرکز معاملات")
    c_buy = m.concentration(conn, "buyer", start, end)
    c_sell = m.concentration(conn, "seller", start, end)
    _write_table(ws9, ["شاخص", "خریداران", "فروشندگان"], [
        ["سهم ۱ نفر برتر (٪)", round(c_buy["top1_share_pct"], 2), round(c_sell["top1_share_pct"], 2)],
        ["سهم ۵ نفر برتر (٪)", round(c_buy["top5_share_pct"], 2), round(c_sell["top5_share_pct"], 2)],
        ["سهم ۱۰ نفر برتر (٪)", round(c_buy["top10_share_pct"], 2), round(c_sell["top10_share_pct"], 2)],
        ["شاخص HHI", round(c_buy["hhi"], 1), round(c_sell["hhi"], 1)],
        ["سطح تمرکز", _level_fa(c_buy["level"]), _level_fa(c_sell["level"])],
    ])

    # --- Sheet 10: Market Maker behavior ---
    ws10 = wb.create_sheet("بازارگردان")
    mm_summary = mm_mod.market_maker_summary(conn, start, end)
    if mm_summary.get("has_market_maker"):
        next_row = _write_narrative(ws10, mm_summary["narrative"], start_row=1, title="رفتار بازارگردان")
        next_row = _write_table(ws10, ["شاخص", "مقدار"], [
            ["بازارگردان(ها)", "، ".join(mm_summary["mm_broker_names"])],
            ["روزهای فعال از کل روزهای معاملاتی", f"{mm_summary['n_active_days']} / {mm_summary['n_trading_days']}"],
            ["روزهای خریدار خالص", mm_summary["n_net_buyer_days"]],
            ["روزهای فروشنده خالص", mm_summary["n_net_seller_days"]],
            ["ارزش کل خرید (ریال)", round(mm_summary["total_buy_value"], 0)],
            ["ارزش کل فروش (ریال)", round(mm_summary["total_sell_value"], 0)],
            ["ارزش خالص (ریال)", round(mm_summary["net_value"], 0)],
            ["تغییر قیمت در بازه (٪)", round(mm_summary["price_change_pct"], 1)],
        ], start_row=next_row)
        daily_rows = [[jalali_date_display(d["date"]), round(d["mm_buy_value"], 0), round(d["mm_sell_value"], 0),
                       round(d["mm_net_value"], 0), round(d["price"], 1)] for d in mm_summary["daily"]]
        _write_table(ws10, ["تاریخ", "خرید بازارگردان (ریال)", "فروش بازارگردان (ریال)", "خالص (ریال)", "قیمت وزنی"],
                     daily_rows, start_row=next_row)
    else:
        _write_narrative(ws10, ["بازارگردانی برای این نماد در این بازه شناسایی نشد."], start_row=1, title="رفتار بازارگردان")

    # --- Sheet 11: Period comparison ---
    ws11 = wb.create_sheet("مقایسه با بازه قبل")
    cmp_res = comp.auto_compare_with_previous(conn, start, end, top_n=top_n)
    header_lines = [
        f"بازه قبل: {cmp_res['period_a']['start_fa']} تا {cmp_res['period_a']['end_fa']}",
        f"بازه فعلی: {cmp_res['period_b']['start_fa']} تا {cmp_res['period_b']['end_fa']}",
    ] + cmp_res["narrative"]
    next_row = _write_narrative(ws11, header_lines, start_row=1, title="مقایسه با بازه قبل")
    delta_rows = []
    label_fa = {"n_transactions": "تعداد معاملات", "total_value": "ارزش کل", "total_quantity": "حجم کل",
                "weighted_avg_price": "میانگین وزنی قیمت", "n_unique_buyers": "تعداد خریداران", "hhi_buyers": "HHI خریداران"}
    for key, d in cmp_res["deltas"].items():
        pct_str = f"{d['pct']:+.1f}٪" if d["pct"] is not None else "—"
        delta_rows.append([label_fa.get(key, key), round(d["a"], 2) if isinstance(d["a"], (int, float)) else d["a"],
                            round(d["b"], 2) if isinstance(d["b"], (int, float)) else d["b"], pct_str])
    next_row2 = _write_table(ws11, ["شاخص", "بازه قبل", "بازه فعلی", "درصد تغییر"], delta_rows, start_row=next_row)
    ws11.cell(row=next_row2, column=1, value=f"خریداران تازه‌وارد: {cmp_res['n_new_participants']} نفر").font = Font(bold=True)
    ws11.cell(row=next_row2 + 1, column=1, value=f"خریداران خارج‌شده: {cmp_res['n_exited_participants']} نفر").font = Font(bold=True)

    # --- Sheet 12: Same-day buy & sell (concrete round-trip evidence) ---
    ws12 = wb.create_sheet("خرید و فروش هم‌روز")
    same_day = bh.same_day_buy_sell_pairs(conn, start, end, top_n=200)
    next_row = _write_narrative(ws12, [bh.DISCLAIMER_SAME_DAY], start_row=1, title="خرید و فروش هم‌روز (شواهد مستقیم نوسان‌گیری)")
    _write_table(ws12, ["نام", "تاریخ", "تعداد خرید", "نرخ خرید", "تعداد فروش", "نرخ فروش",
                         "هم‌پوشانی (٪)", "سود/زیان تقریبی (ریال)", "نوسان‌گیری؟"],
                 [[p["name"], jalali_date_display(p["date"]), p["buy_qty"], p["buy_price"],
                   p["sell_qty"], p["sell_price"], round(p["overlap_ratio"] * 100, 1),
                   p["approx_pnl"], p["verdict"]] for p in same_day],
                 start_row=next_row)

    # --- Sheet 13: Registry / ownership / FIFO ---
    ws13a = wb.create_sheet("مالکیت و FIFO")
    fifo = pf.run_fifo(conn)
    fs = fifo["summary"]
    _write_narrative(ws13a, [
        f"افراد تحلیل‌شده: {fs['people']:,}",
        f"منطبق با رجیستری: {fs['registry_matched']:,}",
        f"موارد بالقوه تازه‌وارد: {fs['potential_new_entrants']:,} — به‌دلیل محدودیت تطبیق هویت، ورود قطعی تلقی نمی‌شود.",
        f"تطبیق کامل پایان دوره: {fs['reconciled']:,} | دارای اختلاف: {fs['mismatched']:,}",
        f"سود تحقق‌یافته با بهای معلوم: {fs['known_realized_pnl']:,.0f} ریال",
        f"فروش مصرف‌کننده موجودی با بهای نامعلوم: {fs['unknown_cost_sold_qty']:,} سهم",
        f"فروش حل‌نشده به‌دلیل کمبود موجودی قابل‌تطبیق: {fs['unresolved_sell_qty']:,} سهم",
        "موجودی اول دوره از رجیستری به‌صورت Lot با بهای تمام‌شده نامعلوم وارد شده است؛ بنابراین سود/زیان آن بخش عددسازی نمی‌شود."
    ], start_row=1, title="خلاصه مالکیت و FIFO")
    fifo_rows=[]
    for r in sorted(fifo["results"], key=lambda x: abs(x["known_realized_pnl"]), reverse=True)[:500]:
        fifo_rows.append([
            r["name"] or "—", r["opening_qty"], r["daily_buy_qty"], r["daily_sell_qty"],
            r["expected_closing_qty"], r["registry_closing_qty"] if r["registry_closing_qty"] is not None else "—",
            r["reconciliation_diff"] if r["reconciliation_diff"] is not None else "—",
            round(r["known_realized_pnl"],0), r["unknown_cost_sold_qty"],
            "بله" if r["is_potential_new_entrant"] else "خیر"
        ])
    _write_table(ws13a,
                 ["نام","اول دوره","خرید روزانه","فروش روزانه","پایان مورد انتظار","پایان رجیستری",
                  "اختلاف","سود تحقق‌یافته معلوم (ریال)","فروش با بهای نامعلوم","ورود احتمالی"],
                 fifo_rows, start_row=12)

    ws13b = wb.create_sheet("تطبیق رجیستری")
    reg_counts = conn.execute("""SELECT
        COUNT(*) AS n,
        SUM(opening_quantity), SUM(registry_buy_quantity), SUM(registry_sell_quantity), SUM(closing_quantity)
        FROM registry_holdings WHERE is_total_row=0""").fetchone()
    ident = conn.execute("""SELECT match_method, COUNT(*) FROM identity_matches GROUP BY match_method""").fetchall()
    _write_table(ws13b, ["شاخص","مقدار"], [
        ["تعداد سهامداران رجیستری", reg_counts[0] or 0],
        ["مجموع اول دوره", reg_counts[1] or 0],
        ["خرید ثبت‌شده در رجیستری", reg_counts[2] or 0],
        ["فروش ثبت‌شده در رجیستری", reg_counts[3] or 0],
        ["مجموع پایان دوره", reg_counts[4] or 0],
        *[[f"تطبیق هویت — {mth}", n] for mth,n in ident],
        ["تطبیق‌نیافته", fs["unmatched_registry"]],
    ])
    _write_narrative(ws13b, [
        "تطبیق موجودی فقط زمانی معنادار است که فایل رجیستری و فایل معاملات دقیقاً به یک بازه و یک وضعیت مالکیت مربوط باشند.",
        "در صورت اختلاف، سامانه اختلاف را پنهان یا با عدد فرضی اصلاح نمی‌کند؛ اختلاف، روش تطبیق و علت احتمالی در گزارش ثبت می‌شود."
    ], start_row=12, title="محدودیت و تفسیر")

    # --- Sheet 13: HHI / Fundamental / Qualitative / Valuation ---
    ws13 = wb.create_sheet("بنیادی-کیفی-ارزش‌گذاری")
    hhi = ext.hhi_summary(conn, start, end)
    row = _write_narrative(ws13, hhi["narrative"], start_row=1, title="۱. شاخص تمرکز (HHI)")
    fund = ext.fundamental_analysis_note(conn)
    row = _write_narrative(ws13, fund["narrative"] + ["--- داده‌های مالی لازم ---"] + fund["required_data"],
                            start_row=row, title="۲. تحلیل بنیادی")
    qual = ext.qualitative_market_behavior(conn, start, end)
    row = _write_narrative(ws13, qual["narrative"] + [
        f"سهم خرید حقوقی: {qual['institutional_share_pct']:.1f}٪",
        f"سهم خرید حقیقی: {qual['individual_share_pct']:.1f}٪",
    ], start_row=row, title="۳. تحلیل کیفی (رفتار بازار)")
    val = ext.valuation_methods_overview(conn, start, end)
    row = _write_narrative(ws13, [f"{me['name']} — داده مورد نیاز: {me['needs']}" for me in val["methods"]] + val["narrative"],
                            start_row=row, title="۴. روش‌های ارزش‌گذاری")

    for ws in wb.worksheets:
        _prepare_sheet(ws)
    pathlib.Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


if __name__ == "__main__":
    db_path = str(pathlib.Path(__file__).resolve().parents[2] / "data" / "smart_equity.db")
    conn = sqlite3.connect(db_path)
    out = generate_excel_report(conn, str(pathlib.Path(__file__).resolve().parents[2] / "exports" / "gozaresh_hiat_modire.xlsx"))
    print("Saved:", out)
